"""
Number-in-Name matcher – all countries in one run (PySpark, distributed).
SQL files only fetch data; all matching logic is here.

Direct record = entity whose name starts with a number.
Good record   = real BvD ID (no '*').
Comparison    = number in Direct name  vs  digits of Good BvD ID, SAME country
                (country = first 2 chars of the ID).

  Step 1  Direct    : number == Good ID digits
  Step 2  Anchored  : one number is the start/end of the other
                      (shorter >= STEP2_MIN_OVERLAP digits, shorter/longer >= STEP2_LENGTH_RATIO_MIN)
  Step 3  Fuzzy     : fuzz.ratio >= DIRECT_THRESHOLD on leading-zero-stripped digits
                      (Direct >= MIN_DIGIT_LENGTH digits, Good length within +-2)
  No name similarity anywhere. A record never matches itself.
  Every join is an equi-join on (Country, key) – no cross joins.
"""
import math
import os
import shutil
import tempfile

from pyspark.sql import DataFrame, Window
from pyspark.sql import functions as F

_STRIP0 = lambda c: F.when(F.regexp_replace(c, "^0+", "") == "", c).otherwise(F.regexp_replace(c, "^0+", ""))


# ─── PREPARE ─────────────────────────────────────────────────────────────────

def prepare_direct(raw: DataFrame, countries=None) -> DataFrame:
    """
    One row per Id. Digits = leading number of the name (ACN/ABN logic removed):
      "123 456 789 ACME"        -> 123456789   (multi-block, used if >= 7 digits)
      "2020 ACME" / "123-456 X" -> 2020 / 123456
      "40624139 - 09082026 APS" -> 40624139    (" - " ends the block)
    """
    df = (raw.groupBy("Id")
             .agg(F.first("Name", ignorenulls=True).alias("Name"),
                  F.max("Information_provider").alias("Information_provider")))
    name = F.trim(F.col("Name"))
    first = F.regexp_replace(F.regexp_extract(name, r"^([0-9]+(-[0-9]+)*)", 1), "[^0-9]", "")
    multi = F.regexp_replace(F.regexp_extract(name, r"^([0-9]+([ -][0-9]+)*)", 1), "[^0-9]", "")
    df = (df.withColumn("Country", F.substring("Id", 1, 2))
            .withColumn("Digits", F.when(F.length(multi) >= 7, multi).otherwise(first)))
    if countries:
        df = df.where(F.col("Country").isin([c.upper() for c in countries]))
    return (df.withColumn("Len", F.length("Digits"))
              .withColumn("Stripped", _STRIP0(F.col("Digits"))))


def prepare_good(raw: DataFrame, direct: DataFrame) -> DataFrame:
    """Good digits = BvD ID without letters. Only countries present in Direct."""
    df = (raw.withColumn("Country", F.substring("Id", 1, 2))
             .withColumn("Digits", F.regexp_replace("Id", "[^0-9]", ""))
             .where(F.col("Digits") != "")
             .join(direct.select("Country").distinct(), "Country", "left_semi"))
    return (df.withColumn("Len", F.length("Digits"))
              .withColumn("Stripped", _STRIP0(F.col("Digits")))
              .select(F.col("Id").alias("Good_ID"), F.col("Name").alias("Good_Name"),
                      "Country", F.col("Digits").alias("G_Digits"),
                      F.col("Len").alias("G_Len"), F.col("Stripped").alias("G_Stripped")))


def _best(cands: DataFrame, order_cols, match_type: str) -> DataFrame:
    """Keep one Good per Direct (first by order_cols) + Candidates = distinct Goods found."""
    w_all = Window.partitionBy("Direct_ID")
    w_ord = Window.partitionBy("Direct_ID").orderBy(*order_cols)
    return (cands.where(F.col("Good_ID") != F.col("Direct_ID"))
                 .withColumn("Candidates", F.size(F.collect_set("Good_ID").over(w_all)))
                 .withColumn("_rn", F.row_number().over(w_ord))
                 .where("_rn = 1").drop("_rn")
                 .withColumn("Match_Type", F.lit(match_type)))


# ─── STEP 1: DIRECT (exact digits) ───────────────────────────────────────────

def step1_exact(direct: DataFrame, good: DataFrame) -> DataFrame:
    d = direct.where(F.col("Digits") != "").select(
        F.col("Id").alias("Direct_ID"), "Country", F.col("Digits").alias("key"))
    g = good.select("Good_ID", "Country", F.col("G_Digits").alias("key"))
    return _best(d.join(g, ["Country", "key"]), [F.col("Good_ID")], "Direct") \
        .select("Direct_ID", "Good_ID", "Candidates", "Match_Type")


# ─── STEP 2: ANCHORED (prefix / suffix, digits only) ─────────────────────────

def _keys(df, digits_col, len_col, min_overlap, ratio_pct):
    """Explode prefix/suffix keys of length k, k in [max(min_overlap, ceil(ratio*len)) .. len-1]."""
    lo = F.greatest(F.lit(min_overlap), F.floor((F.col(len_col) * ratio_pct + 99) / 100).cast("int"))
    hi = F.col(len_col) - 1
    df = df.withColumn("k", F.explode(
        F.when(lo <= hi, F.sequence(lo, hi)).otherwise(F.array().cast("array<int>"))))
    pfx = df.withColumn("key", F.expr(f"substr({digits_col}, 1, k)")).withColumn("Anchor", F.lit("PREFIX"))
    sfx = df.withColumn("key", F.expr(f"substr({digits_col}, {len_col} - k + 1, k)")).withColumn("Anchor", F.lit("SUFFIX"))
    return pfx.unionByName(sfx)


def step2_anchored(rem: DataFrame, good: DataFrame, min_overlap: int, ratio_min: float) -> DataFrame:
    """
    A) Good longer : Good prefixes/suffixes of length = Direct length  == Direct digits
    B) Direct longer: Direct prefixes/suffixes of length = Good length == Good digits
    Best = longest overlap -> smallest length gap -> lowest Good_ID.
    """
    ratio_pct = int(round(ratio_min * 100))
    r = rem.where(F.col("Len") >= min_overlap).select(
        F.col("Id").alias("Direct_ID"), "Country", "Digits", "Len")

    # A) explode Good, only for key lengths that exist in Direct
    d_lens = [row[0] for row in r.select("Len").distinct().collect()]
    gk = _keys(good.select("Good_ID", "Country", "G_Digits", "G_Len"),
               "G_Digits", "G_Len", min_overlap, ratio_pct).where(F.col("k").isin(d_lens))
    cand_a = (r.join(gk, (r.Country == gk.Country) & (r.Len == gk.k) & (r.Digits == gk.key))
               .select("Direct_ID", "Good_ID", "Anchor",
                       r.Len.alias("Overlap"), (gk.G_Len - r.Len).alias("Len_Gap")))

    # B) explode Direct
    dk = _keys(r, "Digits", "Len", min_overlap, ratio_pct)
    g = good.select("Good_ID", "Country", "G_Digits", "G_Len")
    cand_b = (dk.join(g, (dk.Country == g.Country) & (dk.k == g.G_Len) & (dk.key == g.G_Digits))
                .select("Direct_ID", "Good_ID", "Anchor",
                        g.G_Len.alias("Overlap"), (dk.Len - g.G_Len).alias("Len_Gap")))

    return _best(cand_a.unionByName(cand_b),
                 [F.col("Overlap").desc(), F.col("Len_Gap"), F.col("Good_ID")], "Anchored") \
        .select("Direct_ID", "Good_ID", "Candidates", "Match_Type")


# ─── STEP 3: FUZZY DIGITS ────────────────────────────────────────────────────

def _max_indel(l1, l2, thr):
    return math.floor((l1 + l2) * (100.0 - thr) / 100.0 + 1e-9)


def step3_fuzzy(rem: DataFrame, good: DataFrame, threshold: float, min_len: int) -> DataFrame:
    """
    Same result as rapidfuzz cdist(fuzz.ratio, score_cutoff=threshold) on leading-zero-
    stripped digits with Good length bucket within +-2.

    fuzz.ratio = 100 * (1 - indel / (len1 + len2)). At 96:
      * same stripped number                      -> 100
      * one extra / missing digit                 -> only if len1 + len2 >= 25 (13+ digits)
      * one wrong digit (= 2 indels)              -> needs 50+ digits, impossible
    So candidates come from equi-joins: equality + one-digit-deletion keys.
    Best = highest score -> lowest Good_ID.
    """
    if _max_indel(20, 22, threshold) >= 2:
        print(f"WARNING: threshold {threshold} allows 2+ digit edits for long numbers; "
              f"only 0/1-edit matches are searched.")
    del_min = 10**9 if threshold >= 100 else math.ceil((100.0 / (100.0 - threshold) + 1) / 2 - 1e-9)

    r = (rem.where(F.col("Len") >= min_len)
            .withColumn("q", F.regexp_replace("Digits", "^0+", ""))
            .where(F.col("q") != "")
            .select(F.col("Id").alias("Direct_ID"), "Country", "q", F.length("q").alias("Lq")))
    g = good.select("Good_ID", "Country", "G_Len", "G_Stripped", F.length("G_Stripped").alias("Lg"))

    def _del(df, s_col, l_col):
        return (df.where(F.col(l_col) >= del_min)
                  .withColumn("p", F.explode(F.sequence(F.lit(1), F.col(l_col))))
                  .withColumn("dkey", F.expr(f"concat(substr({s_col}, 1, p - 1), substr({s_col}, p + 1))"))
                  .drop("p").dropDuplicates())

    # (1) equal
    c_eq = (r.join(g, (r.Country == g.Country) & (r.q == g.G_Stripped))
             .select("Direct_ID", "Good_ID", "Lq", "Lg", "G_Len"))
    # (2) Direct has one extra digit
    rd = _del(r, "q", "Lq")
    c_dx = (rd.join(g, (rd.Country == g.Country) & (rd.dkey == g.G_Stripped))
              .select("Direct_ID", "Good_ID", "Lq", "Lg", "G_Len"))
    # (3) Good has one extra digit
    gd = _del(g, "G_Stripped", "Lg")
    c_gx = (r.join(gd, (r.Country == gd.Country) & (r.q == gd.dkey))
             .select("Direct_ID", "Good_ID", "Lq", "Lg", "G_Len"))

    cand = (c_eq.unionByName(c_dx).unionByName(c_gx)
              .where(F.abs(F.col("Lq") - F.col("G_Len")) <= 2)
              .withColumn("Score", F.when(F.col("Lq") == F.col("Lg"), F.lit(100.0))
                                    .otherwise(100.0 * (1 - 1.0 / (F.col("Lq") + F.col("Lg")))))
              .where(F.col("Score") >= threshold))
    return _best(cand, [F.col("Score").desc(), F.col("Good_ID")], "Fuzzy") \
        .select("Direct_ID", "Good_ID", "Candidates", "Match_Type", F.round("Score", 2).alias("Score"))


# ─── COMMON LINKS (links_current) ────────────────────────────────────────────

def add_common_links(matches: DataFrame, links: DataFrame, max_ids: int = 10) -> DataFrame:
    """
    SH  = common shareholder  (both are owned by the same entity)
    SUB = common subsidiary   (both own the same entity)
    Direct link (one owns the other) is NOT checked.
    Common_Links: "No" | "SH(2): A, B" | "SH(1): A | SUB(3): X, Y, Z"
    """
    pairs = matches.select("Direct_ID", "Good_ID").distinct()
    ids = pairs.select(F.col("Direct_ID").alias("id")).union(pairs.select(F.col("Good_ID").alias("id"))).distinct()

    lk = links.select("Subsidiary", "Shareholder").distinct()
    by_sub = lk.join(ids, lk.Subsidiary == ids.id, "left_semi")    # rows: id -> its shareholder
    by_sh  = lk.join(ids, lk.Shareholder == ids.id, "left_semi")   # rows: id -> its subsidiary

    def _common(rel, own_col, other_col, label):
        a = rel.select(F.col(own_col).alias("Direct_ID"), F.col(other_col).alias("x"))
        b = rel.select(F.col(own_col).alias("Good_ID"), F.col(other_col).alias("x"))
        return (pairs.join(a, "Direct_ID").join(b, ["Good_ID", "x"])
                     .where((F.col("x") != F.col("Direct_ID")) & (F.col("x") != F.col("Good_ID")))
                     .groupBy("Direct_ID", "Good_ID")
                     .agg(F.sort_array(F.collect_set("x")).alias(label)))

    sh  = _common(by_sub, "Subsidiary", "Shareholder", "sh")
    sub = _common(by_sh, "Shareholder", "Subsidiary", "sub")

    def _fmt(col, tag):
        return F.when(F.col(col).isNotNull(),
                      F.concat(F.lit(f"{tag}("), F.size(col).cast("string"), F.lit("): "),
                               F.array_join(F.slice(col, 1, max_ids), ", ")))

    return (matches.join(sh, ["Direct_ID", "Good_ID"], "left")
                   .join(sub, ["Direct_ID", "Good_ID"], "left")
                   .withColumn("Common_Links",
                               F.when(F.col("sh").isNull() & F.col("sub").isNull(), F.lit("No"))
                                .otherwise(F.concat_ws(" | ", _fmt("sh", "SH"), _fmt("sub", "SUB"))))
                   .drop("sh", "sub"))


# ─── EXCEL OUTPUT ────────────────────────────────────────────────────────────

EXCEL_MAX_ROWS = 1_048_575
MATCH_COLS     = ["Direct_ID", "Direct_Name", "Good_ID", "Good_Name",
                  "Common_Links", "Direct_Information_Provider"]
UNMATCHED_COLS = ["Direct_ID", "Direct_Name", "Extracted_Digits", "Direct_Information_Provider"]
_WIDTHS = {"Direct_Name": 50, "Good_Name": 50, "Common_Links": 60}


def write_excel(sheets, output_path: str):
    """sheets: list of (title, spark_df, cols). Streams rows (no toPandas)."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    def new_sheet(wb, title, cols):
        ws = wb.create_sheet(title[:31])
        for j, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(j)].width = _WIDTHS.get(c, max(len(c) + 4, 14))
        ws.freeze_panes = "A2"
        hdr = []
        for c in cols:
            cell = WriteOnlyCell(ws, value=c.replace("_", " "))
            cell.font = Font(bold=True)
            hdr.append(cell)
        ws.append(hdr)
        return ws

    wb = Workbook(write_only=True)
    for title, df, cols in sheets:
        part, n = 1, 0
        ws = new_sheet(wb, title, cols)
        for row in df.select(*cols).orderBy("Direct_ID").toLocalIterator():
            if n == EXCEL_MAX_ROWS:
                part, n = part + 1, 0
                ws = new_sheet(wb, f"{title} ({part})", cols)
            ws.append(["" if v is None else v for v in row])
            n += 1
        print(f"  {title:<18} {(part - 1) * EXCEL_MAX_ROWS + n:>10,} rows")

    tmp = os.path.join(tempfile.mkdtemp(), os.path.basename(output_path))
    wb.save(tmp)                                   # save locally, then copy (Volume/DBFS safe)
    if output_path.startswith("dbfs:"):
        from pyspark.dbutils import DBUtils
        from pyspark.sql import SparkSession
        DBUtils(SparkSession.getActiveSession()).fs.cp(f"file:{tmp}", output_path)
    else:
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        shutil.copy(tmp, output_path)
    print(f"Saved: {output_path}")
